"""Synthetic portable parser tests. --office uses Desktop's bundled Python."""
import json, subprocess, sys, tempfile, zipfile
from pathlib import Path

root = Path(__file__).resolve().parent.parent
library = Path(sys.argv[sys.argv.index('--library') + 1]) if '--library' in sys.argv else root / 'plugins/dsh-library/lib'
parser = root / 'plugins/dsh-wechat-digest/lib/parse_records.py'
sys.path.insert(0, str(library / 'vendor'))
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

def run(script, path):
    result = subprocess.run([sys.executable, '-I', '-B', str(script), str(path)], capture_output=True, text=True)
    value = json.loads(result.stdout)
    assert result.returncode in (0, 1) and (not value['ok'] or result.returncode == 0)
    return value

with tempfile.TemporaryDirectory(prefix='dsh-parsers-') as temp:
    folder = Path(temp)
    pdf = folder / 'lesson.pdf'
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=400)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    fonts = DictionaryObject({NameObject('/F1'): font})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): fonts})
    content = DecodedStreamObject(); content.set_data(b'BT /F1 14 Tf 20 300 Td (Classroom PDF sample) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(content)
    with pdf.open('wb') as stream: writer.write(stream)
    result = run(library / 'reader.py', pdf)
    assert result['ok'] and 'Classroom PDF sample' in result['value']['text']
    messages = folder / 'synthetic.json'
    messages.write_text(json.dumps([{'sender':'Student A','timestamp':'2026-10-15 09:00','text':'Please review the demo.'}, {'sender':'Student B','timestamp':'2026-10-15 09:01','text':'I will review it today.'}]))
    result = run(parser, messages)
    assert result['ok'] and result['value']['recordCount'] == 2
    bad = folder / 'external.xml'
    bad.write_text('<!DOCTYPE x [<!ENTITY external SYSTEM "file:///not-a-real-file">]><x>&external;</x>')
    assert run(parser, bad)['ok'] is False
    archive = folder / 'outside.zip'
    with zipfile.ZipFile(archive, 'w') as z: z.writestr('../outside.txt', 'must not be imported')
    assert run(parser, archive)['ok'] is False
    if '--office' in sys.argv:
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        d = Document(); d.add_paragraph('Classroom DOCX sample'); d.save(folder / 'lesson.docx')
        w = Workbook(); w.active['A1'] = 'Classroom XLSX sample'; w.save(folder / 'lesson.xlsx')
        p = Presentation(); s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = 'Classroom PPTX sample'; p.save(folder / 'lesson.pptx')
        for suffix in ['docx','xlsx','pptx']:
            result = run(library / 'reader.py', folder / ('lesson.' + suffix))
            assert result['ok'] and ('Classroom ' + suffix.upper() + ' sample') in result['value']['text']
print('PASS synthetic PDF, chat records, XML entity denial and ZIP path boundary' + ('; bundled Office parsers' if '--office' in sys.argv else ''))
