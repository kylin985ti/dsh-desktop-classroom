"""Read only the explicitly selected document using Desktop's bundled parsers."""
import sys, json, zipfile, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'vendor'))
LIMIT = 300000

def guard_archive(path):
    if not zipfile.is_zipfile(path):
        raise ValueError('办公文件格式无效')
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 10000 or sum(info.file_size for info in infos) > 64 * 1024 * 1024:
            raise ValueError('办公文件解压后超过读取上限')
        if any(info.flag_bits & 1 or info.file_size > 1024 * 1024 and info.file_size / max(info.compress_size, 1) > 300 for info in infos):
            raise ValueError('办公文件的压缩格式不适合读取')

def read_document(path):
    extension = os.path.splitext(path)[1].lower()
    lines = []
    warnings = ['这是正文提取；版式和批注请用原应用查看。']
    truncated = False
    count = 0
    def add(value):
        nonlocal truncated, count
        if count + len(value) + 1 > LIMIT:
            truncated = True
            return False
        lines.append(value)
        count += len(value) + 1
        return True
    if extension == '.pdf':
        from pypdf import PdfReader
        document = PdfReader(path)
        if document.is_encrypted:
            raise ValueError('PDF 已加密，请先在原应用解锁')
        for number, page in enumerate(document.pages[:500], 1):
            if not add('第 %s 页\n%s' % (number, page.extract_text() or '')):
                break
        if len(document.pages) > 500:
            truncated = True
    elif extension == '.docx':
        guard_archive(path)
        from docx import Document
        document = Document(path)
        for paragraph in document.paragraphs:
            if not add(paragraph.text):
                break
        for table in document.tables:
            for row in table.rows:
                if not add(' | '.join(cell.text for cell in row.cells)):
                    break
    elif extension in ('.xlsx', '.xlsm'):
        guard_archive(path)
        from openpyxl import load_workbook
        document = load_workbook(path, read_only=True, data_only=True)
        warnings.append('公式显示文件中保存的计算结果，不会执行宏或重新计算。')
        try:
            for sheet in document:
                add('工作表：' + sheet.title)
                for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if index > 20000 or not add(' | '.join('' if cell is None else str(cell) for cell in row)):
                        truncated = True
                        break
        finally:
            document.close()
    elif extension == '.pptx':
        guard_archive(path)
        from pptx import Presentation
        document = Presentation(path)
        for number, slide in enumerate(document.slides, 1):
            add('第 %s 张幻灯片' % number)
            for shape in slide.shapes:
                if shape.has_text_frame:
                    add(shape.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        add(' | '.join(cell.text for cell in row.cells))
    else:
        raise ValueError('这种格式请用对应应用打开')
    text = '\n'.join(lines).strip()
    if not text:
        warnings.append('未提取到正文；扫描件可先用原应用进行文字识别。')
    if truncated:
        warnings.append('文件较长，当前正文提取到达上限。')
    return {'text': text, 'warnings': warnings, 'truncated': truncated}

try:
    print(json.dumps({'ok': True, 'value': read_document(sys.argv[1])}, ensure_ascii=False))
except Exception as error:
    message = str(error) if isinstance(error, ValueError) else '无法提取文件正文，请用对应应用打开检查'
    print(json.dumps({'ok': False, 'error': message}, ensure_ascii=False))
